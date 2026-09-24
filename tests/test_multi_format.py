#!/usr/bin/env python3
# 多角度跨格式症状测试：docx/xlsx/pptx/eml 各自的相对位置对齐是否成立。
# 核心验证点（用户工作流关键）：
#   1) 内容编辑/对调 => 0 偏差（不误报）
#   2) 插入/删除行·列·幻灯片 => 无级联（removed/added 孤立，不牵连后续）
#   3) 多 part（多 sheet/多 slide）同内容 => 0 偏差（不撞车）
#   4) 格式偏离 => 报 changed
import os, sys, json, tempfile
SKILL = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, SKILL)
import checker as C
import openpyxl
from openpyxl.styles import Font
from pptx import Presentation
from pptx.util import Pt

TMP = tempfile.mkdtemp(prefix="otc_mf_")
cfg = json.load(open(os.path.join(SKILL, "param_tree.json"), encoding="utf-8"))
pt = C.ParamTree(cfg)


def compare(t, d):
    devs, red = C.compare(t, d, cfg, pt)
    from collections import Counter
    return devs, Counter(x['kind'] for x in devs), red


results = []
def check(name, ok, detail=""):
    results.append((name, ok, detail))
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f"  -- {detail}" if detail else ""))


# ---------------- docx（沿用既有 build_e2e 的样本构造）----------------
import build_e2e_samples  # 触发样本生成并跑一次（其内已 print 报告）

# ---------------- xlsx ----------------
def mk_xlsx(path, data, fonts=None):
    wb = openpyxl.Workbook()
    names = list(data.keys())
    wb.active.title = names[0]
    for n in names[1:]:
        wb.create_sheet(n)
    for name, rows in data.items():
        ws = wb[name]
        for r, row in enumerate(rows, 1):
            for c, val in enumerate(row, 1):
                ws.cell(r, c, val)
    if fonts:
        for (sname, r, c), f in fonts.items():
            wb[sname].cell(r, c).font = f
    wb.save(path)


mk_xlsx(f"{TMP}/x_t.xlsx", {"S1": [["制造商","深圳"],["型号","ABC-123"],["容量","3000mAh"]]})
mk_xlsx(f"{TMP}/x_ins.xlsx", {"S1": [["制造商","深圳"],["型号","ABC-123"],["容量","3000mAh"],["生产日期","2026-09"]]})
mk_xlsx(f"{TMP}/x_del.xlsx", {"S1": [["制造商","深圳"],["容量","3000mAh"]]})  # 删中间行
mk_xlsx(f"{TMP}/x_edit.xlsx", {"S1": [["制造商","深圳"],["型号","ABC-123X"],["容量","3000mAh"]]})  # 内容编辑
mk_xlsx(f"{TMP}/x_fmt.xlsx", {"S1": [["制造商","深圳"],["型号","ABC-123"],["容量","3000mAh"]]},
        fonts={("S1",2,2): Font(size=20)})  # 格式偏离
mk_xlsx(f"{TMP}/x_swap.xlsx", {"S1": [["型号","ABC-123"],["制造商","深圳"],["容量","3000mAh"]]})  # 内容对调
mk_xlsx(f"{TMP}/x_2s_t.xlsx", {"S1":[["a","1"]], "S2":[["b","2"]]})
mk_xlsx(f"{TMP}/x_2s_d.xlsx", {"S1":[["a","1"]], "S2":[["b","2"]]})  # 多 sheet 同内容

d, k, _ = compare(f"{TMP}/x_t.xlsx", f"{TMP}/x_ins.xlsx")
check("xlsx 插入行(末): 无级联且新增", k['removed']==0 and k['added']>0, f"k={dict(k)}")
d, k, _ = compare(f"{TMP}/x_t.xlsx", f"{TMP}/x_del.xlsx")
check("xlsx 删除中间行: 无级联(added=0)", k['removed']>0 and k['added']==0, f"k={dict(k)}")
d, k, _ = compare(f"{TMP}/x_t.xlsx", f"{TMP}/x_edit.xlsx")
check("xlsx 单元格内容编辑 => 0 偏差", len(d)==0, f"n={len(d)}")
d, k, _ = compare(f"{TMP}/x_t.xlsx", f"{TMP}/x_fmt.xlsx")
check("xlsx 格式偏离 => changed>0", k['changed']>0 and k['removed']==0, f"k={dict(k)}")
d, k, _ = compare(f"{TMP}/x_t.xlsx", f"{TMP}/x_swap.xlsx")
check("xlsx 内容对调 => 0 偏差", len(d)==0, f"n={len(d)}")
d, k, _ = compare(f"{TMP}/x_2s_t.xlsx", f"{TMP}/x_2s_d.xlsx")
check("xlsx 多sheet同内容 => 0 偏差(不撞车)", len(d)==0, f"n={len(d)}")


# ---------------- pptx ----------------
def mk_pptx(path, slides):
    prs = Presentation()
    for shapes in slides:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        for (txt, sz) in shapes:
            tb = s.shapes.add_textbox(0,0,400,100)
            run = tb.text_frame.paragraphs[0].add_run(); run.text = txt
            if sz: run.font.size = Pt(sz)
    prs.save(path)

mk_pptx(f"{TMP}/p_t.pptx", [[("标题",18)], [("副标",14)]])
mk_pptx(f"{TMP}/p_ins.pptx", [[("标题",18)], [("副标",14)], [("新页",12)]])  # 插幻灯片
mk_pptx(f"{TMP}/p_del.pptx", [[("标题",18)]])  # 删幻灯片
mk_pptx(f"{TMP}/p_edit.pptx", [[("标题改",18)], [("副标",14)]])  # 内容编辑
mk_pptx(f"{TMP}/p_fmt.pptx", [[("标题",40)], [("副标",14)]])  # 改字号
mk_pptx(f"{TMP}/p_2s_t.pptx", [[("A",18)], [("B",14)]])
mk_pptx(f"{TMP}/p_2s_d.pptx", [[("A",18)], [("B",14)]])  # 多 slide 同内容

d, k, _ = compare(f"{TMP}/p_t.pptx", f"{TMP}/p_ins.pptx")
check("pptx 插幻灯片: 无级联且新增", k['removed']==0 and k['added']>0, f"k={dict(k)}")
d, k, _ = compare(f"{TMP}/p_t.pptx", f"{TMP}/p_del.pptx")
check("pptx 删幻灯片: removed>0 added=0", k['removed']>0 and k['added']==0, f"k={dict(k)}")
d, k, _ = compare(f"{TMP}/p_t.pptx", f"{TMP}/p_edit.pptx")
check("pptx 文本框内容编辑 => 0 偏差", len(d)==0, f"n={len(d)}")
d, k, _ = compare(f"{TMP}/p_t.pptx", f"{TMP}/p_fmt.pptx")
check("pptx 改文本框字号 => changed>0", k['changed']>0 and k['removed']==0, f"k={dict(k)}")
d, k, _ = compare(f"{TMP}/p_2s_t.pptx", f"{TMP}/p_2s_d.pptx")
check("pptx 多slide同内容 => 0 偏差(不撞车)", len(d)==0, f"n={len(d)}")


# ---------------- eml ----------------
def mk_eml(path, subject, styled_lines, content_type="text/html; charset=utf-8"):
    # styled_lines: [(text, style_str), ...]
    header = (f"From: a@x.com\r\nTo: b@x.com\r\nSubject: {subject}\r\n"
              f"Content-Type: {content_type}\r\n\r\n")
    body = "<html><body>" + "".join(
        f'<p style="{st}">{txt}</p>' for txt, st in styled_lines
    ) + "</body></html>"
    with open(path, "w", encoding="utf-8", newline="") as f:
        f.write(header + body)

SAME = "color:#000000;font-size:14px"
mk_eml(f"{TMP}/e_t.eml", "检测报告",
       [("制造商：深圳市某公司",SAME),("型号：ABC-123",SAME),("容量：3000mAh",SAME)])
# 改头 + 改样式颜色 + 插段落（型号用 span 改红）
mk_eml(f"{TMP}/e_style.eml", "检测报告（修订）",
       [("制造商：深圳市某公司",SAME),('<span style="color:#FF0000">型号：ABC-123X</span>',"color:#FF0000"),
        ("容量：3000mAh",SAME),("生产日期：2026-09",SAME)])
# 仅插段落（内容编辑，样式不变）
mk_eml(f"{TMP}/e_ins.eml", "检测报告",
       [("制造商：深圳市某公司",SAME),("型号：ABC-123",SAME),("容量：3000mAh",SAME),("备注：无",SAME)])
# 仅内容编辑（文字变，样式完全相同）
mk_eml(f"{TMP}/e_edit.eml", "检测报告",
       [("制造商：深圳市某某公司",SAME),("型号：ABC-123",SAME),("容量：3000mAh",SAME)])

d, k, _ = compare(f"{TMP}/e_t.eml", f"{TMP}/e_style.eml")
subj_changed = any(x['element']=='Subject' and x['kind']=='changed' for x in d)
check("eml 改头+改样式+插段落: 报Subject改 & 有新增", subj_changed and k['added']>0, f"k={dict(k)}")
d, k, _ = compare(f"{TMP}/e_t.eml", f"{TMP}/e_ins.eml")
check("eml 仅插段落(内容编辑): 有新增无级联", k['added']>0 and k['removed']==0, f"k={dict(k)}")
d, k, _ = compare(f"{TMP}/e_t.eml", f"{TMP}/e_edit.eml")
check("eml 仅内容编辑(样式同) => 0 偏差", len(d)==0, f"n={len(d)}")


# ---------------- 汇总 ----------------
print("\n==== 汇总 ====")
np = sum(1 for _, ok, _ in results if ok)
print(f"{np}/{len(results)} 通过")
fails = [n for n, ok, _ in results if not ok]
if fails:
    print("失败项:", fails)
    raise SystemExit(1)
print("ALL PASS")
